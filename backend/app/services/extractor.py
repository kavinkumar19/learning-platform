"""
Text extraction service for PDF, DOCX, PPTX and TXT files.
No AI or ML is used — only conventional parsing libraries.
"""
from pathlib import Path


def extract_text(file_path: str, file_type: str) -> str:
    """
    Extract readable text from a file.
    Returns extracted text as a plain string.
    Raises RuntimeError on unsupported file type or extraction failure.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if file_type == "pdf":
        return _extract_pdf(path)
    elif file_type == "docx":
        return _extract_docx(path)
    elif file_type == "pptx":
        return _extract_pptx(path)
    elif file_type == "txt":
        return path.read_text(encoding="utf-8", errors="replace")
    else:
        raise RuntimeError(f"Unsupported file type: {file_type}")


def _extract_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        pages = []
        for page in doc:
            pages.append(page.get_text())
        doc.close()
        text = "\n".join(pages).strip()
        if not text:
            raise RuntimeError(
                "No text could be extracted from this PDF. "
                "It may be a scanned image PDF; OCR is not supported."
            )
        return text
    except ImportError:
        raise RuntimeError("PyMuPDF (fitz) is not installed")
    except Exception as e:
        raise RuntimeError(f"PDF extraction failed: {e}")


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except ImportError:
        raise RuntimeError("python-docx is not installed")
    except Exception as e:
        raise RuntimeError(f"DOCX extraction failed: {e}")


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        slides = []
        for slide in prs.slides:
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides.append("\n".join(texts))
        return "\n\n".join(slides)
    except ImportError:
        raise RuntimeError("python-pptx is not installed")
    except Exception as e:
        raise RuntimeError(f"PPTX extraction failed: {e}")
